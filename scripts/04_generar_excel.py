"""Genera un dashboard en Excel a partir de la BD SQLite."""
from pathlib import Path
import sqlite3
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils.dataframe import dataframe_to_rows

BASE = Path(__file__).resolve().parent.parent
DB = BASE / "data" / "construccion.db"
OUT = BASE / "excel" / "dashboard_construccion.xlsx"


def _hoja(wb, nombre, df):
    ws = wb.create_sheet(nombre)
    for r in dataframe_to_rows(df, index=False, header=True):
        ws.append(r)
    return ws


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(DB)
    prov = pd.read_sql(
        "SELECT d.provincia, COUNT(*) num_permisos "
        "FROM permisos p JOIN dim_provincia d ON p.provincia_id=d.provincia_id "
        "GROUP BY d.provincia ORDER BY num_permisos DESC LIMIT 10", con)
    evol = pd.read_sql(
        "SELECT anio, COUNT(*) num_permisos, "
        "ROUND(SUM(valor_edificacion_usd), 0) valor_total_usd "
        "FROM permisos GROUP BY anio ORDER BY anio", con)
    mat = pd.read_sql(
        "SELECT material_estructura, COUNT(*) num_permisos "
        "FROM permisos WHERE material_estructura IS NOT NULL "
        "GROUP BY material_estructura ORDER BY num_permisos DESC", con)
    con.close()

    wb = Workbook()
    wb.remove(wb.active)

    ws1 = _hoja(wb, "Top provincias", prov)
    chart = BarChart()
    chart.title = "Top 10 provincias por permisos"
    chart.type = "bar"
    data = Reference(ws1, min_col=2, min_row=1, max_row=len(prov) + 1)
    cats = Reference(ws1, min_col=1, min_row=2, max_row=len(prov) + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height, chart.width = 9, 16
    ws1.add_chart(chart, "E2")

    ws2 = _hoja(wb, "Evolucion anual", evol)
    line = LineChart()
    line.title = "Evolucion de permisos por anio"
    d2 = Reference(ws2, min_col=2, min_row=1, max_row=len(evol) + 1)
    c2 = Reference(ws2, min_col=1, min_row=2, max_row=len(evol) + 1)
    line.add_data(d2, titles_from_data=True)
    line.set_categories(c2)
    line.height, line.width = 9, 16
    ws2.add_chart(line, "E2")

    ws3 = _hoja(wb, "Materiales", mat)
    bar3 = BarChart()
    bar3.title = "Permisos por material de estructura"
    d3 = Reference(ws3, min_col=2, min_row=1, max_row=len(mat) + 1)
    c3 = Reference(ws3, min_col=1, min_row=2, max_row=len(mat) + 1)
    bar3.add_data(d3, titles_from_data=True)
    bar3.set_categories(c3)
    bar3.height, bar3.width = 9, 16
    ws3.add_chart(bar3, "E2")

    wb.save(OUT)
    print(f"[ok] Excel generado en {OUT}")


if __name__ == "__main__":
    main()
